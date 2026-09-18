"""Arizalarni Excel faylga chiqarish."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from bot.db.models import Ariza, Fan, HOLAT_NOMI, Holat, User
from bot.utils.vaqt import mahalliy

SARLAVHALAR = [
    "№",
    "Ariza raqami",
    "Familiya",
    "Ism",
    "Sharif",
    "Telefon",
    "Maktab",
    "Sinf",
    "Fan",
    "Holat",
    "Ro'yxatdan o'tgan",
    "Telegram",
]

KENGLIKLAR = [5, 18, 16, 14, 18, 16, 34, 6, 20, 18, 18, 16]


async def arizalar_excel(
    session: AsyncSession,
    *,
    boshlanish: datetime | None = None,
    tugash: datetime | None = None,
    fan_id: int | None = None,
    sinf: int | None = None,
    holat: Holat | None = None,
) -> tuple[BytesIO, int]:
    """Excel faylni va undagi qatorlar sonini qaytaradi.

    `boshlanish`/`tugash` — ariza ochilgan vaqt bo'yicha filtr (UTC da beriladi).
    """
    sorov = (
        select(Ariza, User, Fan)
        .join(User, Ariza.user_id == User.id)
        .join(Fan, Ariza.fan_id == Fan.id)
        .where(Ariza.holat != Holat.BEKOR_QILINGAN)
        .order_by(Fan.tartib, User.sinf, User.familiya)
    )
    if boshlanish is not None:
        sorov = sorov.where(Ariza.created_at >= boshlanish)
    if tugash is not None:
        sorov = sorov.where(Ariza.created_at < tugash)
    if fan_id is not None:
        sorov = sorov.where(Ariza.fan_id == fan_id)
    if sinf is not None:
        sorov = sorov.where(User.sinf == sinf)
    if holat is not None:
        sorov = sorov.where(Ariza.holat == holat)

    qatorlar = (await session.execute(sorov)).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Arizalar"

    sarlavha_fill = PatternFill("solid", start_color="1F4E78")
    sarlavha_font = Font(bold=True, color="FFFFFF")

    ws.append(SARLAVHALAR)
    for ustun in range(1, len(SARLAVHALAR) + 1):
        katak = ws.cell(row=1, column=ustun)
        katak.fill = sarlavha_fill
        katak.font = sarlavha_font
        katak.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(ustun)].width = KENGLIKLAR[ustun - 1]

    for nomer, (ariza, user, fan) in enumerate(qatorlar, start=1):
        ws.append(
            [
                nomer,
                ariza.ariza_raqami or "",
                user.familiya,
                user.ism,
                user.sharif,
                user.telefon,
                user.maktab,
                user.sinf,
                fan.nomi,
                HOLAT_NOMI[ariza.holat],
                mahalliy(ariza.created_at),
                f"@{user.username}" if user.username else "",
            ]
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SARLAVHALAR))}{ws.max_row}"

    fayl = BytesIO()
    wb.save(fayl)
    fayl.seek(0)
    return fayl, len(qatorlar)
